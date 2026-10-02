"""Test export uploaded from the UI (#207): POST/GET/DELETE /api/testset, lookup order, predict.csv, 413/422."""

from __future__ import annotations

import base64
import csv
import io
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api import app as app_mod
from src.api import template as template_mod
from src.api import testset as T
from src.ingest.predict import COLUMNS

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "testset"
NOTICES, ITEMS = "Извещения_24-25.csv", "ТРУ_24-25.csv"


def _text(name: str) -> str:
    return (FIX / name).read_text(encoding="utf-8")


def _files(**extra: str) -> list[dict[str, str]]:
    """The fixture export as the UI sends it: file name + text; `extra` adds or replaces files."""
    files = {NOTICES: _text(NOTICES), ITEMS: _text(ITEMS), **extra}
    return [{"name": n, "text": t} for n, t in files.items()]


@pytest.fixture
def root(tmp_path: Path) -> Path:
    return tmp_path / "testsets"


@pytest.fixture
def client(root: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setenv("TESTSETS_DIR", str(root))
    monkeypatch.delenv("TEST_DATA_DIR", raising=False)
    T.reset()
    return TestClient(app_mod.app)


def _upload(client: TestClient, files: list[dict[str, str]]):
    return client.post("/api/testset", json={"files": files})


def test_upload_lists_lots_and_becomes_active(client, root):
    r = _upload(client, _files(**{"Поставщики_24-25.csv": "lot_id;supplier_inn\n9001;7801000012\n"}))
    assert r.status_code == 200, r.text
    body = r.json()
    assert re.fullmatch(r"[0-9a-f]{12}", body["id"])
    assert body["files"] == [NOTICES, ITEMS] and body["lots"] == 3
    assert [it["lot_id"] for it in body["items"]] == ["9001", "9002", "9003"]
    first = body["items"][0]
    assert first == {"lot_id": "9001", "reqnum": "0172200000126000001", "text": "Поставка бумаги офисной А4",
                     "okpd2": ["17.12.14.110"], "notes": []}
    assert body["items"][1]["reqnum"] is None
    # Поставщики*.csv: a warning, never stored or read (the answer must not come from it)
    assert any("Поставщики_24-25.csv" in w and "не используется" in w for w in body["warnings"])
    folder = root / body["id"]
    assert sorted(p.name for p in folder.glob("*.csv")) == ["Извещения.csv", "ТРУ.csv"]
    assert not any("Поставщики" in p.name for p in root.rglob("*"))
    assert (root / "active").read_text(encoding="utf-8").strip() == body["id"]
    assert client.get("/api/testset").json() == body


def test_id_is_the_content_hash_without_suppliers(client):
    a = _upload(client, _files()).json()["id"]
    b = _upload(client, _files(**{"Поставщики.csv": "anything"})).json()["id"]
    assert a == b  # the suppliers file does not take part: same export, same id
    c = _upload(client, _files(**{ITEMS: _text(ITEMS).replace("17.12.14.110", "17.12.14.120")})).json()["id"]
    assert c != a
    assert client.get("/api/testset").json()["id"] == c  # the last upload is active


def test_get_and_delete(client, root):
    assert client.get("/api/testset").status_code == 404
    assert client.delete("/api/testset").status_code == 204  # nothing active: still fine
    body = _upload(client, _files()).json()
    assert client.get("/api/testset").status_code == 200
    r = client.delete("/api/testset")
    assert r.status_code == 204 and r.content == b""
    assert client.get("/api/testset").status_code == 404
    assert not (root / "active").exists()
    assert (root / body["id"] / "Извещения.csv").exists()  # the files stay in the cache


def test_active_survives_restart(client, search_data):
    body = _upload(client, _files()).json()
    T.reset()  # a new API process: nothing in memory, only cache/testsets/active
    assert client.get("/api/testset").json() == body
    assert client.get("/api/lot/9003").json()["okpd2"] == ["26.70.22.150"]


def test_items_without_okpd2_carry_notes(client):
    body = _upload(client, [{"name": NOTICES, "text": _text(NOTICES)}]).json()
    assert any("ТРУ" in w for w in body["warnings"])
    assert all(it["okpd2"] == [] and it["notes"] for it in body["items"])


def test_lot_lookup_uploaded_then_test_dir_then_data(client, search_data, tmp_path, monkeypatch):
    test_dir = tmp_path / "test-dir"
    test_dir.mkdir()
    (test_dir / NOTICES).write_text("lot_id;subject\n9001;Из TEST_DATA_DIR\n9;Вода из TEST_DATA_DIR\n",
                                    encoding="utf-8")
    monkeypatch.setenv("TEST_DATA_DIR", str(test_dir))
    assert client.get("/api/lot/9001").json()["text"] == "Из TEST_DATA_DIR"
    notices = _text(NOTICES) + "2026-01-18;904;9;1.00;;Вода;Вода из загрузки;false;;;ЭМ\n"
    assert _upload(client, _files(**{NOTICES: notices})).status_code == 200
    assert client.get("/api/lot/9001").json()["text"] == "Поставка бумаги офисной А4"  # uploaded first
    assert client.get("/api/lot/0172200000126000001").json()["query_id"] == "lot-9001"
    assert client.get("/api/lot/9").json()["text"] == "Вода из загрузки"  # over TEST_DATA_DIR and the data
    client.delete("/api/testset")
    assert client.get("/api/lot/9001").json()["text"] == "Из TEST_DATA_DIR"
    monkeypatch.delenv("TEST_DATA_DIR")
    assert client.get("/api/lot/9").json()["text"] == "Поставка воды питьевой"  # training data
    assert client.get("/api/lot/9001").status_code == 404


def _cp1251_as_browser_text() -> str:
    # the UI reads a file with File.text(): bytes that are not UTF-8 become U+FFFD
    return "lot_id;subject\n1;Бумага\n".encode("cp1251").decode("utf-8", errors="replace")


@pytest.mark.parametrize(
    ("files", "needle"),
    [
        ([{"name": "Лоты.csv", "text": "x"}], "Лоты.csv"),
        ([], "Извещения"),
        ([{"name": ITEMS, "text": _text(ITEMS)}], "Извещения"),
        ([{"name": NOTICES, "text": _text(NOTICES)}, {"name": "Извещения_2.csv", "text": _text(NOTICES)}],
         "несколько"),
        ([{"name": "Извещения.xlsx", "text": "PK\x03\x04"}], "CSV"),
        ([{"name": NOTICES, "text": "lot_id;name\n1;x\n"}], "Извещения_24-25.csv: нет колонок subject"),
        ([{"name": NOTICES, "text": "lot_id,subject\n1,x\n"}], "разделитель"),
        ([{"name": NOTICES, "text": ""}], "пустой файл"),
        ([{"name": NOTICES, "text": _cp1251_as_browser_text()}], "UTF-8"),
        ([{"name": NOTICES, "text": _text(NOTICES)}, {"name": ITEMS, "text": "lot_id;product_name\n1;x\n"}],
         "ТРУ_24-25.csv: нет колонок okpd2_code"),
    ],
)
def test_format_errors_are_422_with_a_message(client, root, files, needle):
    before = _upload(client, _files()).json()
    r = _upload(client, files)
    assert r.status_code == 422
    detail = r.json()["detail"]
    assert isinstance(detail, str) and needle in detail
    assert str(root) not in detail and "testsets" not in detail  # no server path in the answer
    assert client.get("/api/testset").json()["id"] == before["id"]  # the active set is unchanged
    assert sorted(p.name for p in root.iterdir()) == sorted([before["id"], "active"])  # nothing left behind


def test_defense_export_names_are_accepted(client, root):
    # the organizers' defense files (#596): «Предзащита_Извещения_N» + «Предзащита_Потоварка_N» (ТРУ)
    files = [{"name": "Предзащита_Извещения_4.csv", "text": _text(NOTICES)},
             {"name": "Предзащита_Потоварка_4.csv", "text": _text(ITEMS)},
             {"name": "Предзащита_Поставщики_4.csv", "text": "lot_id;supplier_inn\n9001;7801000012\n"}]
    r = _upload(client, files)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["files"] == ["Предзащита_Извещения_4.csv", "Предзащита_Потоварка_4.csv"] and body["lots"] == 3
    assert body["items"][0]["okpd2"] == ["17.12.14.110"]
    assert any("Предзащита_Поставщики_4.csv" in w and "не используется" in w for w in body["warnings"])
    assert sorted(p.name for p in (root / body["id"]).glob("*.csv")) == ["Извещения.csv", "ТРУ.csv"]
    assert _upload(client, [{"name": "Структура.csv", "text": "x"}]).status_code == 422  # «тру» inside a word
    assert _upload(client, [{"name": "Извещения1.csv", "text": _text(NOTICES)}]).status_code == 200  # a word's start


def test_folder_name_from_the_browser_is_only_a_name(client, root):
    files = [{"name": "../../Извещения_24-25.csv", "text": _text(NOTICES)},
             {"name": "C:\\fakepath\\ТРУ_24-25.csv", "text": _text(ITEMS)}]
    body = _upload(client, files).json()
    assert body["files"] == [NOTICES, ITEMS] and body["lots"] == 3
    assert {p.parent for p in root.rglob("*.csv")} == {root / body["id"]}


def test_upload_size_limit_is_413(client, monkeypatch):
    size = len(_text(NOTICES).encode("utf-8")) + len(_text(ITEMS).encode("utf-8"))
    monkeypatch.setattr(T, "MAX_UPLOAD_BYTES", size)
    assert _upload(client, _files()).status_code == 200  # exactly at the limit
    monkeypatch.setattr(T, "MAX_UPLOAD_BYTES", size - 1)
    r = _upload(client, _files())
    assert r.status_code == 413 and "МБ" in r.json()["detail"]


def test_oversized_body_is_413_before_parsing(client, monkeypatch):
    monkeypatch.setattr(T, "MAX_UPLOAD_BYTES", 1000)
    junk = b"x" * (T.max_body_bytes() + 1)  # not even JSON: rejected by size, not by the parser
    r = client.post("/api/testset", content=junk, headers={"content-type": "application/json"})
    assert r.status_code == 413
    r = client.post("/api/testset", content=b"x" * 10, headers={"content-type": "application/json"})
    assert r.status_code == 422  # a small broken body reaches the parser


def test_predict_csv(client, search_data, root):
    assert client.get("/api/testset/predict.csv").status_code == 404
    set_id = _upload(client, _files()).json()["id"]
    r = client.get("/api/testset/predict.csv?top=3&offline=1")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "text/csv; charset=utf-8"
    assert r.headers["content-disposition"] == f'attachment; filename="mitrandir-predict-{set_id}.csv"'
    assert r.content.startswith("\ufeff".encode())  # BOM: Excel opens UTF-8
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert tuple(rows[0]) == COLUMNS
    lot = COLUMNS.index("lot_id")
    assert {row[lot] for row in rows[1:]} == {"9001", "9002", "9003"}
    rank, block = COLUMNS.index("rank"), COLUMNS.index("block")
    assert max(int(row[rank]) for row in rows[1:] if row[block] == "История") <= 3
    again = client.get("/api/testset/predict.csv?top=3&offline=1")
    assert again.content == r.content  # idempotent: the finished file is served again
    one = client.get("/api/testset/predict.csv?top=1&offline=1").content.decode("utf-8-sig")
    hist = [row for row in csv.reader(io.StringIO(one), delimiter=";") if row[block] == "История"]
    assert hist and all(row[rank] == "1" for row in hist)  # another top is another file, not a resume
    assert len(list((root / set_id).glob("predict*.csv"))) == 2


def test_predict_csv_without_training_data_is_503(client, tmp_path, monkeypatch):
    from src.match import index

    _upload(client, _files())
    index.clear_memory_cache()
    monkeypatch.setenv("DATA_STORE", "parquet")
    monkeypatch.setenv("DATA_PROCESSED_DIR", str(tmp_path / "no-processed"))
    monkeypatch.setenv("INDEX_DIR", str(tmp_path / "indexes"))
    r = client.get("/api/testset/predict.csv")
    assert r.status_code == 503
    assert "no-processed" not in r.json()["detail"]


@pytest.mark.parametrize("query", ["top=0", "top=51", "offline=maybe"])
def test_predict_csv_rejects_bad_params(client, query):
    _upload(client, _files())
    assert client.get(f"/api/testset/predict.csv?{query}").status_code == 422


def _predict_files(root: Path, set_id: str) -> list[Path]:
    return sorted((root / set_id).glob("predict-*.csv"))


def test_upload_again_recomputes_the_prediction(client, search_data, root):
    """Review #216: the same content uploaded again (new data, new code) must not serve the old file."""
    set_id = _upload(client, _files()).json()["id"]
    first = client.get("/api/testset/predict.csv?top=3").content
    (path,) = _predict_files(root, set_id)
    with path.open("a", encoding="utf-8", newline="") as f:
        f.write("stale;row\n")  # what an older run left behind
    assert b"stale;row" in client.get("/api/testset/predict.csv?top=3").content  # served from the file
    assert _upload(client, _files()).json()["id"] == set_id
    again = client.get("/api/testset/predict.csv?top=3").content
    assert b"stale;row" not in again and again == first


def test_another_scoring_config_is_another_file(client, search_data, root, tmp_path, monkeypatch):
    from src.match import search as search_mod

    set_id = _upload(client, _files()).json()["id"]
    client.get("/api/testset/predict.csv?top=3")
    (old,) = _predict_files(root, set_id)
    assert old.name == f"predict-{T.config_tag()}-top3.csv"
    cfg = tmp_path / "search.yaml"
    cfg.write_bytes(search_mod.SEARCH_CONFIG.read_bytes() + b"\n# other weights\n")
    monkeypatch.setattr(search_mod, "SEARCH_CONFIG", cfg)
    assert client.get("/api/testset/predict.csv?top=3").status_code == 200
    names = {p.name for p in _predict_files(root, set_id)}
    assert names == {old.name, f"predict-{T.config_tag()}-top3.csv"} and len(names) == 2


def test_live_run_over_every_lot_is_refused(client):
    _upload(client, _files())
    r = client.get("/api/testset/predict.csv?offline=0")
    assert r.status_code == 422 and "массовый сбор" in r.json()["detail"]


def test_predict_csv_makes_no_request_even_with_discovery_online(client, search_data, monkeypatch):
    """offline=1 holds whatever DISCOVERY_OFFLINE says: zero requests to the open sources."""
    import httpx

    calls: list[str] = []

    def count(self: httpx.HTTPTransport, request: httpx.Request) -> httpx.Response:
        calls.append(str(request.url))
        return httpx.Response(503)

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", count)
    monkeypatch.setenv("DISCOVERY_OFFLINE", "0")
    monkeypatch.delenv("ENRICH_OFFLINE", raising=False)
    _upload(client, _files())
    assert client.get("/api/testset/predict.csv?top=3&offline=1").status_code == 200
    assert calls == []


# --- the customer's list of purchases by the template (#501) ---

def _template_body(name: str, data: bytes) -> dict[str, str]:
    return {"name": name, "data": base64.b64encode(data).decode("ascii")}


def test_template_download(client):
    r = client.get("/api/testset/template.xlsx")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="mitrandir-template.xlsx"'
    wb = load_workbook(io.BytesIO(r.content))
    assert [c.value for c in wb["Закупки"][1]] == list(template_mod.HEADERS) and "Инструкция" in wb.sheetnames


def test_template_upload_then_excel_with_the_errors_sheet(client, search_data, root):
    text = ("Номер закупки;Предмет закупки;ОКПД2;ИНН заказчика;НМЦК, руб.\n"
            "A-1;Поставка бумаги офисной А4;17.12.14.110;;\n"
            "A-2;Ремонт кровли;;;\n"
            "A-3;;;;\n")
    r = client.post("/api/testset/template", json=_template_body("Мои закупки.csv", text.encode("cp1251")))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["lots"] == 2 and body["files"] == ["Мои закупки.csv"]
    assert [(it["lot_id"], it["reqnum"]) for it in body["items"]] == [("2", "A-1"), ("3", "A-2")]
    assert any("не принято строк — 1" in w for w in body["warnings"])
    assert any("строка 4" in w and "искать не по чему" in w for w in body["warnings"])
    assert client.get("/api/lot/3").json()["text"] == "Ремонт кровли"  # the UI picks by lot_id: the row

    x = client.get("/api/testset/predict.xlsx?top=3")
    assert x.status_code == 200, x.text
    assert x.headers["content-disposition"] == f'attachment; filename="mitrandir-predict-{body["id"]}.xlsx"'
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Рекомендации", "Ошибки"]
    rec = list(wb["Рекомендации"].iter_rows(values_only=True))
    assert tuple(rec[0]) == COLUMNS
    assert len(rec) > 1 and {str(row[COLUMNS.index("lot_id")]) for row in rec[1:]} <= {"2", "3"}
    errors = [tuple("" if v is None else str(v) for v in row) for row in wb["Ошибки"].iter_rows(values_only=True)]
    assert errors[0] == T.ERROR_HEADERS
    assert ("4", "A-3") in {(e[0], e[1]) for e in errors[1:] if e[3] == "строка не принята"}
    # the lot without OKPD2 is searched, with its note
    assert any(e[0] == "3" and "ОКПД2 не указан" in e[2] and e[3] == "подобрано с ограничением" for e in errors[1:])


def test_template_upload_errors(client):
    assert client.post("/api/testset/template", json={"name": "x.xlsx", "data": "@@@"}).status_code == 422
    r = client.post("/api/testset/template", json=_template_body("x.csv", "Цвет\nсиний\n".encode()))
    assert r.status_code == 422 and "Предмет закупки" in r.json()["detail"]
    assert client.get("/api/testset").status_code == 404  # nothing became active


def test_predict_xlsx_needs_a_set(client):
    assert client.get("/api/testset/predict.xlsx").status_code == 404


def test_template_row_number_wins_over_a_users_number(client):
    # the user's «Номер закупки» 3 is row 2; row 3 has its own subject: a pick by lot_id must open row 3
    text = "Номер закупки;Предмет закупки\n3;Бумага\n7;Ремонт кровли\n"
    assert client.post("/api/testset/template", json=_template_body("l.csv", text.encode())).status_code == 200
    assert client.get("/api/lot/3").json()["text"] == "Ремонт кровли"
    assert client.get("/api/lot/7").json()["text"] == "Ремонт кровли"  # the user's number still finds its row


def test_errors_sheet_is_text_not_formulas(client, search_data):
    text = "Номер закупки;Предмет закупки;ИНН заказчика\n=1+1;Бумага\x02 офисная;123\n"
    assert client.post("/api/testset/template", json=_template_body("l.csv", text.encode())).status_code == 200
    x = client.get("/api/testset/predict.xlsx?top=3")
    assert x.status_code == 200, x.text  # a control character from the data does not break the book
    err = load_workbook(io.BytesIO(x.content))["Ошибки"]
    assert all(c.data_type != "f" for row in err.iter_rows() for c in row)
