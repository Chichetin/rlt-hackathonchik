"""The customer's list of purchases (#501): template, xlsx / cp1251 CSV, synonyms, rows that do not pass."""

from __future__ import annotations

import csv
import io

import pytest
from openpyxl import Workbook, load_workbook

from src.api import template as T


def _valid_inn() -> str:
    for tail in range(10):
        inn = f"780234567{tail}"
        if T._inn_ok(inn):
            return inn
    raise AssertionError("no valid INN")


def _xlsx(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = T.SHEET
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _notices(p: T.Parsed) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(p.notices), delimiter=";"))


def _items(p: T.Parsed) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(p.items), delimiter=";"))


def test_template_has_the_columns_and_the_instruction():
    wb = load_workbook(io.BytesIO(T.template()))
    assert wb.sheetnames == [T.SHEET, "Инструкция"]
    assert [c.value for c in wb[T.SHEET][1]] == list(T.HEADERS)
    assert wb[T.SHEET]["C2"].number_format == "@"  # OKPD2 stays text in Excel
    # the template's own header parses back into every field
    assert {T._BY_NAME[T._norm(h)] for h in T.HEADERS} == set(T.FIELDS)


def test_xlsx_rows_become_the_organizers_csv():
    inn = _valid_inn()
    p = T.parse("закупки.xlsx", _xlsx([
        list(T.HEADERS),
        ["0172200000126000001", "Поставка бумаги А4", "17.12.14.110, 17.23.13", inn, 125000.5],
        [None, "Ремонт кровли", None, None, None],
    ]))
    assert p.lots == 2 and p.errors == []
    n = _notices(p)
    assert [r["lot_id"] for r in n] == ["2", "3"]  # the row of the file
    assert n[0]["reqnum"] == "0172200000126000001" and n[0]["customer_inn"] == inn
    assert n[0]["start_price"] == "125000.5" and n[0]["subject"] == "Поставка бумаги А4"
    assert [(r["lot_id"], r["okpd2_code"]) for r in _items(p)] == [("2", "17.12.14.110"), ("2", "17.23.13")]


def test_cp1251_csv_with_russian_synonyms():
    text = "Реестровый номер;Наименование;Код ОКПД2;ИНН;НМЦК\n42;Канцтовары;;;1 500,00\n"
    p = T.parse("список.csv", text.encode("cp1251"))
    assert p.lots == 1 and p.errors == []
    (row,) = _notices(p)
    assert (row["reqnum"], row["subject"], row["start_price"]) == ("42", "Канцтовары", "1500.00")
    assert _items(p) == []  # empty OKPD2: searched by the subject (load_testset adds its own note)


def test_utf8_csv_with_commas():
    p = T.parse("list.csv", "﻿предмет,окпд2\nБумага,17.12\n".encode())
    assert p.lots == 1 and _items(p)[0]["okpd2_code"] == "17.12"


def test_rows_that_do_not_pass_and_notes():
    p = T.parse("закупки.xlsx", _xlsx([
        [None],  # a blank row above the header: row numbers stay the user's
        list(T.HEADERS),
        ["1", "Бумага", "abc", "1234567890", "дорого"],
        ["2", None, None, None, None],
        ["1", "Ещё бумага", "17.12", None, None],
        [None, None, None, None, None],
    ]))
    assert p.lots == 1
    got = [(e.row, e.reqnum, e.kept) for e in p.errors]
    assert got == [(3, "1", True), (3, "1", True), (3, "1", True), (4, "2", False), (5, "1", False)]
    reasons = " | ".join(e.reason for e in p.errors)
    assert "ОКПД2 «abc»" in reasons and "ИНН заказчика «1234567890»" in reasons and "НМЦК «дорого»" in reasons
    assert "искать не по чему" in reasons and "номер повторяется (уже в строке 3)" in reasons
    (row,) = _notices(p)
    assert row["customer_inn"] == "" and row["start_price"] == ""


@pytest.mark.parametrize(("name", "data", "needle"), [
    ("x.pdf", b"%PDF", "нужен .xlsx или .csv"),
    ("x.xlsx", b"not a zip", "не читается как .xlsx"),
    ("x.csv", b"", "пустой файл"),
    ("x.csv", "Цвет;Размер\nсиний;1\n".encode(), "нет колонок «Предмет закупки» и «ОКПД2»"),
    ("x.csv", "Предмет;ОКПД2\n;\n;\n".encode(), "ни одной закупки"),
])
def test_unusable_files_say_what_to_do(name, data, needle):
    with pytest.raises(T.TemplateError, match=needle):
        T.parse(name, data)


def test_inn_control_digits():
    assert T._inn_ok("7707083893")  # Сбербанк
    assert not T._inn_ok("7707083894")
    assert T._inn_ok("500100732259")  # a valid 12-digit one
    assert not T._inn_ok("12345")
