"""API: the contract the UI and the e2e demo are built against; search on the ingest fixture."""

import io
import json
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api import app as app_mod
from src.contracts import Counterparty, CounterpartyMatch, Evidence, ScoreComponent, inn_is_valid


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{(tmp_path / 'q.sqlite').as_posix()}")
    monkeypatch.setattr(app_mod, "_engine", None)
    # no ingest unless a test builds one (_lots): /api/lot reads the store only (#293)
    monkeypatch.setenv("DATA_PROCESSED_DIR", str(tmp_path / "no-processed"))
    monkeypatch.setattr(app_mod, "_probe", app_mod.Probe())  # the health probe result is cached per process (#492)
    return TestClient(app_mod.app)


@pytest.fixture
def search_client(client, search_data):
    """client whose search reads the ingest fixture (DATA_PROCESSED_DIR from search_data)."""
    return client


def test_health_reports_no_stub(search_client):
    body = search_client.get("/api/health").json()
    assert body["status"] == "ok" and body["stub"] is False and body["llm"] == "none"


def test_health_is_not_ok_when_search_fails(search_client, monkeypatch, caplog):
    """#492: a green health while every search answers 500 is a dead demo that looks alive."""
    from src.match import search as search_mod

    def broken(*a, **kw):
        raise RuntimeError("index is broken")

    monkeypatch.setattr(search_mod, "search", broken)
    assert search_client.get("/api/health").json()["status"] == "search_failed"
    assert "index is broken" in caplog.text


def test_health_is_not_ok_without_data(client):
    assert client.get("/api/health").json()["status"] == "no_data"


def test_health_reuses_the_probe_result(search_client, monkeypatch):
    from src.match import search as search_mod

    calls = []
    real = search_mod.search
    monkeypatch.setattr(search_mod, "search", lambda *a, **kw: calls.append(1) or real(*a, **kw))
    assert search_client.get("/api/health").json()["status"] == "ok"
    assert search_client.get("/api/health").json()["status"] == "ok"
    assert len(calls) == 1


def test_health_probe_is_offline(search_client, monkeypatch):
    """The probe is the /api/search path without «Новые» (registries) and without the enrichment queue."""
    from src.match import pipeline

    seen = []
    real = pipeline.pick
    monkeypatch.setattr(pipeline, "pick", lambda q, **kw: seen.append(kw) or real(q, **kw))
    assert search_client.get("/api/health").json()["status"] == "ok"
    assert seen == [{"discovery": False}]


def test_health_times_out_on_a_hung_search(search_client, monkeypatch):
    import threading

    from src.match import search as search_mod

    release = threading.Event()
    monkeypatch.setattr(app_mod, "PROBE_TIMEOUT", 0.05)
    monkeypatch.setattr(search_mod, "search", lambda *a, **kw: release.wait(5) and [])
    try:
        assert search_client.get("/api/health").json()["status"] == "search_timeout"
    finally:
        release.set()
        for t in threading.enumerate():  # the probe ends before the fixtures are torn down
            if t.name == "health-probe":
                t.join(5)


def test_search_two_blocks_real_inns(search_client):
    body = search_client.post("/api/search", json={"query_id": "q1", "text": "бумага А4"}).json()
    assert body["history"] and body["stub"] is False
    assert body["discovered"] == []  # open-source block is #16
    assert all(m["origin"] == "history" for m in body["history"])
    assert all(c["explanation"] for m in body["history"] for c in m["components"])
    inns = [m["counterparty"]["inn"] for m in body["history"] if m["counterparty"]["inn"]]
    assert inns[0] == "7801000012" and all(inn_is_valid(i) for i in inns)


def test_search_without_data_is_503(client):
    r = client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert r.status_code == 503 and "src.ingest.build" in r.json()["detail"]
    assert "no-processed" not in r.json()["detail"]  # no server path in the answer
    assert client.post("/api/export", json={"query_id": "q", "text": "бумага"}).status_code == 503


def test_search_rejects_empty_text(client):
    assert client.post("/api/search", json={"query_id": "q", "text": ""}).status_code == 422


def test_search_rejects_client_as_of_without_building_an_index(search_client, search_data):
    """as_of is the backtest knob (eval calls search() directly). Through the API a client date built a new BM25
    index for that date under the global index lock: on the full data 46 s and 110 MB per date, every other
    search waiting (40 s instead of 0.5 s), then 503 (#145, docs/REVIEW.md F-01)."""
    import os
    from pathlib import Path

    for path in ("/api/search", "/api/export"):
        r = search_client.post(path, json={"query_id": "q", "text": "бумага", "as_of": "2025-03-01"})
        assert r.status_code == 422 and "as_of" in r.json()["detail"]
    import json

    built = Path(os.environ["INDEX_DIR"])
    # nothing was built for the client's date; the full index (as_of None) is the product's own — #589 reads its
    # data_end to accept a lot after the data (#605)
    dirs = built.glob("bm25-*") if built.exists() else []
    metas = [json.loads((d / "meta.json").read_text(encoding="utf-8")) for d in dirs]
    assert all(m.get("as_of") is None for m in metas), metas



def test_search_from_a_notice_says_the_notice_is_left_out(search_client):
    """#249, F-13: the UI explains why the loaded notice is not among its own results (X-Excluded-Lot)."""
    r = search_client.post("/api/search", json={"query_id": "lot-6", "text": "Бумага офисная А4"})
    assert r.status_code == 200 and r.headers["X-Excluded-Lot"] == "6"
    assert all(e["url"] != "data:items#6" for m in r.json()["history"] for e in m["counterparty"]["evidence"])
    free = search_client.post("/api/search", json={"query_id": "ui-1", "text": "Бумага офисная А4"})
    assert "X-Excluded-Lot" not in free.headers

def test_search_rejects_oversized_text(search_client):
    """Unbounded text held a worker: 1 MB — 18 s on the full data; the longest eval query is 4.5k chars (#145, F-02)."""
    big = "бумага офисная " * 1400  # 21 000 chars
    for path in ("/api/search", "/api/export"):
        r = search_client.post(path, json={"query_id": "q", "text": big})
        assert r.status_code == 422 and "символов" in r.json()["detail"]
    fits = search_client.post("/api/search", json={"query_id": "q", "text": "бумага офисная " * 1300})  # 19 500
    assert fits.status_code == 200 and fits.json()["history"]


def test_search_pages_on_server_with_totals(search_client):
    q = {"query_id": "q", "text": "бумага"}
    full = search_client.post("/api/search", json=q)
    total = int(full.headers["X-Total-History"])
    rows = [m["counterparty"]["inn"] for m in full.json()["history"]]
    assert total == len(rows) >= 3 and full.headers["X-Total-Discovered"] == "0"
    page = search_client.post("/api/search?limit=2&offset=1", json=q)
    assert page.headers["X-Total-History"] == str(total)  # the total does not depend on the page
    assert [m["counterparty"]["inn"] for m in page.json()["history"]] == rows[1:3]
    assert search_client.post(f"/api/search?offset={total}", json=q).json()["history"] == []
    assert search_client.post("/api/search?limit=0", json=q).status_code == 422
    assert search_client.post("/api/search?offset=-1", json=q).status_code == 422


def test_filters_apply_before_the_page(search_client, monkeypatch):
    """A filter selects from every candidate, not from the first page: its total is of the filtered rows."""
    from src.match import search as search_mod
    seen = []
    real = search_mod.search
    monkeypatch.setattr(search_mod, "search", lambda q, k=20: seen.append(k) or real(q, k=k))
    q = {"query_id": "q", "text": "бумага", "filters": {"region": "Санкт-Петербург"}}
    r = search_client.post("/api/search?limit=1", json=q)
    assert seen == [search_mod.CANDIDATES]
    spb = int(r.headers["X-Total-History"])
    assert spb >= 2 and len(r.json()["history"]) == 1
    assert r.json()["history"][0]["counterparty"]["kpp"].startswith("78")


def test_health_names_the_store(client):
    assert client.get("/api/health").json()["store"] in ("parquet", "postgres")


def test_search_filters(search_client):
    # kind/МСП are not known from organizers' data (enrich): a filter on them drops, not guesses
    body = search_client.post("/api/search", json={"query_id": "q", "text": "бумага",
                                                   "filters": {"only_manufacturers": True}}).json()
    assert body["history"] == [] and body["discovered"] == []


def test_search_region_filter_by_kpp(search_client):
    body = search_client.post("/api/search", json={"query_id": "q", "text": "бумага",
                                                   "filters": {"region": "Санкт-Петербург"}}).json()
    kpps = [m["counterparty"]["kpp"] for m in body["history"]]
    assert kpps and all(k.startswith("78") for k in kpps)
    lo = search_client.post("/api/search", json={"query_id": "q", "text": "бумага",
                                                 "filters": {"region": "Ленинградская область"}}).json()
    assert [m["counterparty"]["inn"] for m in lo["history"]] == ["4701000038"]


def test_counterparty(client, tmp_path, monkeypatch):
    # real enrich (pb.nalog.ru via cache) — offline and with an empty cache an INN is not found;
    # the card itself is covered in tests/unit/test_enrich.py
    monkeypatch.setenv("ENRICH_OFFLINE", "1")
    monkeypatch.setenv("SOURCE_CACHE", str(tmp_path / "sources.sqlite"))
    assert client.get("/api/counterparty/0000100012").status_code == 404
    assert client.get("/api/counterparty/7707083894").status_code == 422


def test_enqueue(client):
    r = client.post("/api/enrich/0000100012")
    assert r.status_code == 202 and r.json()["task_id"] >= 1


def test_reliability_is_read_for_the_shown_page_and_the_export_rows_only(search_client, monkeypatch):
    """#127 review: shadow reliability is one history SQL over the page (≤ limit INNs) — not over the ≤ 200
    candidates the pipeline scores; the export reads it for its own rows."""
    from src.match import search as search_mod
    from src.match.store import ParquetStore

    cfg = search_mod.load_config()
    monkeypatch.setattr(search_mod, "load_config", lambda: {**cfg, "reliability": {"mode": "shadow"}})  # on since #585
    calls: list[int] = []
    real = ParquetStore.supplier_history
    monkeypatch.setattr(ParquetStore, "supplier_history", lambda self, inns, as_of: calls.append(len(inns))
                        or real(self, inns, as_of))
    q = {"query_id": "q", "text": "бумага офисная"}
    r = search_client.post("/api/search?limit=2", json=q)
    page = r.json()["history"]
    assert int(r.headers["X-Total-History"]) > 2 and calls == [len(page)] == [2]
    rel = [c for m in page for c in m["components"] if c["name"] == "reliability"]
    assert all(c["weight"] == 0 and c["explanation"].endswith("справочно, в оценку не входит") for c in rel)
    calls.clear()
    search_client.post("/api/export", json=q)
    assert calls == [int(r.headers["X-Total-History"])]  # the exported rows, once


def test_export_xlsx(search_client):
    r = search_client.post("/api/export", json={"query_id": "q", "text": "бумага"})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    header = [c.value for c in ws[1]]
    assert header[:3] == ["Блок", "Наименование", "ИНН"] and ws.max_row > 1
    assert ws.cell(row=2, column=3).value == "7801000012"
    # the order and the relevance are two columns (#278, D-156): the place is the block's order, not the score's
    place, score = header.index("Место в блоке") + 1, header.index("Релевантность (формула критериев)") + 1
    assert "Выше следующего" in header and "Score" not in header
    history = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "С опытом госзакупок"]
    assert [ws.cell(row=r, column=place).value for r in history] == list(range(1, len(history) + 1))
    assert all(isinstance(ws.cell(row=r, column=score).value, float) for r in history)


def test_export_kind_column_in_russian(client, monkeypatch):
    """«Роль» of the xlsx: the captions of the UI (KIND_LABEL in src/ui/src/api.ts), not the contract's keys (#102,
    #232: the column and «не определена» are called «Роль», as the chips)."""
    from src.match import search as search_mod
    kinds = ["manufacturer", "distributor", "dealer", "unknown", None, "not-a-kind"]  # last: a code from the future
    matches = []
    for i, kind in enumerate(kinds):
        m = _fabricated("history")
        matches.append(m.model_copy(update={"counterparty": m.counterparty.model_copy(
            update={"kind": kind, "name": f"ООО Тип {i}"})}))
    monkeypatch.setattr(search_mod, "IS_STUB", True)  # no cache / discovery: only the column is under test
    monkeypatch.setattr(search_mod, "search", lambda q, k=20: matches)
    r = client.post("/api/export", json={"query_id": "q", "text": "x"})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    assert [c.value for c in ws[1]][3] == "Роль"
    assert [ws.cell(row=row, column=4).value for row in range(2, ws.max_row + 1)] == [
        "Производитель", "Дистрибьютор", "Дилер", "Роль не определена", "Роль не определена", "Роль не определена"]


def test_openapi_has_contracts(client):
    schemas = client.get("/openapi.json").json()["components"]["schemas"]
    assert {"Query", "Evidence", "Counterparty", "ScoreComponent", "CounterpartyMatch"} <= set(schemas)


def test_unknown_api_path_is_404_not_spa(client, tmp_path, monkeypatch):
    assert client.get("/api/nope").status_code == 404


def _fabricated(origin: str) -> CounterpartyMatch:
    ts = datetime(2026, 10, 1, tzinfo=UTC)
    cp = Counterparty(
        inn="7801000012", name="ИНН 7801000012", contacts={"email": "fake@example.org"},
        evidence=[Evidence(source="data:participants", url="data:participants#1", fetched_at=ts, fact="ИНН",
                           field="inn", value="7801000012"),
                  Evidence(source="сайт компании", url="https://example.org/x", fetched_at=ts, fact="email",
                           field="email", value="fake@example.org")])
    return CounterpartyMatch(counterparty=cp, score=0.5, origin=origin,  # type: ignore[arg-type]
                             components=[ScoreComponent(name="product_fit", value=0.5, weight=0.35, explanation="x")])


def test_search_strips_unverified_facts(client, monkeypatch, tmp_path):
    """A fabricated contact never reaches the card or the xlsx."""
    from src.match import search as search_mod
    monkeypatch.setattr(search_mod, "search", lambda q, k=20: [_fabricated("history"), _fabricated("discovered")])
    from src.match import pipeline
    monkeypatch.setattr(pipeline, "SourceCache", lambda: __import__("src.sources.cache", fromlist=["x"]).SourceCache(
        tmp_path / "c.sqlite"))
    body = client.post("/api/search", json={"query_id": "q", "text": "x"}).json()
    rows = body["history"] + body["discovered"]
    # web evidence (example.org) is not in the empty cache -> dropped, contacts dropped
    assert len(rows) == 2 and all(m["counterparty"]["contacts"] == {} for m in rows)
    assert all(e["url"].startswith("data:") for m in rows for e in m["counterparty"]["evidence"])
    assert all(m["counterparty"]["inn"] == "7801000012" for m in rows)


def _lots(tmp_path, monkeypatch, notices: list[dict], items: list[dict]) -> None:
    """An ingest directory with notices and items only: /api/lot reads it through the store (#293)."""
    import pyarrow as pa
    import pyarrow.parquet as pq

    n = pa.schema([("lot_id", pa.int64()), ("procedure_id", pa.int64()), ("reqnum", pa.string()),
                   ("procedure_name", pa.string()), ("subject", pa.string()), ("start_price", pa.float64()), ("customer_inn", pa.string()),
                   ("channel", pa.string())])
    i = pa.schema([("lot_id", pa.int64()), ("product_name", pa.string()), ("okpd2_code", pa.string())])
    pq.write_table(pa.Table.from_pylist([{f: r.get(f) for f in n.names} for r in notices], schema=n),
                   tmp_path / "notices.parquet")
    pq.write_table(pa.Table.from_pylist(items, schema=i), tmp_path / "items.parquet")
    monkeypatch.setenv("DATA_PROCESSED_DIR", str(tmp_path))


DEMO_LOTS = [
    {"lot_id": 9000001, "procedure_name": "Бумага", "subject": "Бумага офисная А4, 80 г/м², 500 листов",
     "start_price": 45000.0, "customer_inn": "7804105239", "channel": "ЭМ"},
    {"lot_id": 9000002, "procedure_id": 1841706, "reqnum": "0372200260025000001", "procedure_name": "Услуги",
     "subject": "Оказание услуг связи", "start_price": 120000.0, "customer_inn": "7804105239", "channel": "АИС ГЗ"},
]
DEMO_ITEMS = [{"lot_id": 9000001, "product_name": "Бумага для офисной техники", "okpd2_code": "17.12.14.119"},
              {"lot_id": 9000001, "product_name": "Бумага", "okpd2_code": "17.12.14.110"},
              {"lot_id": 9000002, "product_name": "Услуги связи", "okpd2_code": "61.10.11.110"}]


def test_lot_by_lot_id_and_reqnum(client, monkeypatch, tmp_path):
    _lots(tmp_path, monkeypatch, DEMO_LOTS, DEMO_ITEMS)
    q = client.get("/api/lot/9000001").json()
    assert q["query_id"] == "lot-9000001" and q["text"].startswith("Бумага офисная А4")
    assert q["budget"] == 45000.0 and q["customer_inn"] == "7804105239" and q["channel"] == "ЭМ"
    assert q["okpd2"] == ["17.12.14.110", "17.12.14.119"]
    by_req = client.get("/api/lot/0372200260025000001").json()
    assert by_req["query_id"] == "lot-9000002" and by_req["channel"] == "АИС ГЗ"


def test_procedure_opens_its_lot(client, monkeypatch, tmp_path):
    """#617: the procedure number (Извещения.procedure_id) gives the same query as the number of its lot."""
    _lots(tmp_path, monkeypatch, DEMO_LOTS, DEMO_ITEMS)
    assert client.get("/api/procedure/1841706").json() == client.get("/api/lot/9000002").json()
    assert client.get("/api/procedure/9000001").status_code == 404  # a lot_id is not a procedure number
    assert client.get("/api/procedure/" + "9" * 30).status_code == 404  # longer than bigint: not a crash
    assert client.get("/api/procedure/abc").status_code == 422


def test_lot_not_found_404(client, monkeypatch, tmp_path):
    _lots(tmp_path, monkeypatch, DEMO_LOTS, DEMO_ITEMS)
    assert client.get("/api/lot/1").status_code == 404


def test_lot_without_data_503(client):
    """No ingest built (the client's DATA_PROCESSED_DIR is empty): 503 with the command, not a raw-CSV guess."""
    r = client.get("/api/lot/9000001")
    assert r.status_code == 503 and "данные" in r.json()["detail"] and "src.ingest.build" in r.json()["detail"]


def test_lot_rejects_non_digits(client):
    assert client.get("/api/lot/abc").status_code == 422


def test_lot_broken_parquet_503_not_crash(client, monkeypatch, tmp_path):
    _lots(tmp_path, monkeypatch, DEMO_LOTS, DEMO_ITEMS)
    (tmp_path / "notices.parquet").write_bytes(b"not a parquet file")
    assert client.get("/api/lot/9000001").status_code == 503


def test_lot_without_price_gives_no_budget(client, monkeypatch, tmp_path):
    # ingest TRY_CASTs start_price: «n/a» in the CSV is NULL here (src/ingest/build.py)
    _lots(tmp_path, monkeypatch, [{"lot_id": 5, "subject": "Товар", "channel": "ЭМ"}], [])
    q = client.get("/api/lot/5").json()
    assert q["budget"] is None and q["customer_inn"] is None and q["okpd2"] == []


@pytest.mark.parametrize(("procedure", "text"), [
    ("Поставка картриджей", "Поставка картриджей"),
    ("Закупка, осуществляемая в соответствии с частью 12 статьи 93", "Картридж; Тонер"),  # #113: both templates
])
def test_lot_template_subject_falls_back_to_procedure_name(client, monkeypatch, tmp_path, procedure, text):
    _lots(tmp_path, monkeypatch,
          [{"lot_id": 7, "procedure_name": procedure, "start_price": 1000.0, "channel": "АИС ГЗ",
            "subject": "Закупка у единственного поставщика частью 12 статьи 93"}],
          [{"lot_id": 7, "product_name": "Тонер", "okpd2_code": "20.59"},
           {"lot_id": 7, "product_name": "Картридж", "okpd2_code": "26.20"}])
    assert client.get("/api/lot/7").json()["text"] == text


def test_lot_reads_ingest_parquet(client, monkeypatch, tmp_path):
    import duckdb
    con = duckdb.connect()
    con.execute("COPY (SELECT 77::BIGINT lot_id, '0372200260024000124' reqnum, 'Картриджи' procedure_name, "
                "'Картриджи' subject, 500.0 start_price, '7804105239' customer_inn, 'ЭМ' channel) "
                f"TO '{(tmp_path / 'notices.parquet').as_posix()}' (FORMAT parquet)")
    con.execute("COPY (SELECT 77::BIGINT lot_id, 'x' product_name, '26.20.40.190' okpd2_code) "
                f"TO '{(tmp_path / 'items.parquet').as_posix()}' (FORMAT parquet)")
    monkeypatch.setenv("DATA_PROCESSED_DIR", str(tmp_path))
    q = client.get("/api/lot/0372200260024000124").json()
    assert q["query_id"] == "lot-77" and q["okpd2"] == ["26.20.40.190"]
    assert q["channel"] == "ЭМ" and q["budget"] == 500.0


def test_health_stub_flag_follows_demo_fixture_env(client, monkeypatch):
    """scripts/demo.sh without organizers' data sets DEMO_FIXTURE=1 → the UI shows the demo banner."""
    monkeypatch.delenv("DEMO_FIXTURE", raising=False)
    assert client.get("/api/health").json()["stub"] is False
    monkeypatch.setenv("DEMO_FIXTURE", "1")
    assert client.get("/api/health").json()["stub"] is True


def test_okpd2_labels_for_ui(client):
    r = client.get("/api/okpd2", params={"codes": "17.12.14.110, 80.10.12.000,99.99.99.999,,"})
    assert r.status_code == 200
    assert r.json() == {"17.12.14.110": "17.12.14.110 «Бумага для печати»",
                        "80.10.12.000": "80.10.12.000 (по 80.10.12 «Услуги охраны»)",
                        "99.99.99.999": "99.99.99.999"}
    assert client.get("/api/okpd2").json() == {}
    assert client.get("/api/okpd2", params={"codes": ",".join(f"17.12.{i}" for i in range(51))}).status_code == 422


# --- #232: role chips (kinds), ЮЛ/ИП (entity), minimum relevance (min_score), X-Facets ---------------------------

ROLE_ROWS = [  # name, kind, inn (10 digits — ЮЛ, 12 — ИП), score, status, in_rnp
    ("Завод", "manufacturer", "7801000012", 0.9, None, None),
    ("Опт", "distributor", "780100000090", 0.7, None, None),
    ("Дилер", "dealer", "4701000038", 0.5, None, None),
    ("Ремонт", "unknown", "470100000088", 0.3, None, None),
    ("Без роли", None, None, 0.2, None, None),
    ("Ликвидирован", "manufacturer", "7801000020", 0.95, "В процессе ликвидации", None),
    ("В РНП", "distributor", "4703173460", 0.85, None, True),
]
ZERO = {"kind": {"manufacturer": 0, "distributor": 0, "dealer": 0, "unknown": 0},
        "entity": {"ul": 0, "ip": 0, "unknown": 0}, "near": {"worked": 0, "near": 0, "any": 0}}


def _role_rows() -> list[CounterpartyMatch]:
    return [CounterpartyMatch(
        counterparty=Counterparty(inn=inn, name=name, kind=kind, status=status, in_rnp=rnp),  # type: ignore[arg-type]
        score=score, origin="history",
        components=[ScoreComponent(name="product_fit", value=score, weight=1, explanation="x")])
        for name, kind, inn, score, status, rnp in ROLE_ROWS]


@pytest.fixture
def roles_client(client, monkeypatch):
    """The search answers ROLE_ROWS; stub path — no cache, no discovery: only the filters are under test."""
    from src.match import search as search_mod
    monkeypatch.setattr(search_mod, "IS_STUB", True)
    monkeypatch.setattr(search_mod, "search", lambda q, k=20: _role_rows())
    return client


def _names(client, **filters) -> list[str]:
    r = client.post("/api/search", json={"query_id": "q", "text": "x", "filters": filters})
    assert r.status_code == 200, r.text
    names = [m["counterparty"]["name"] for m in r.json()["history"]]
    assert r.headers["X-Total-History"] == str(len(names))
    return names


def test_search_kinds_chips(roles_client):
    assert _names(roles_client) == ["Завод", "Опт", "Дилер", "Ремонт", "Без роли"]  # liquidated / RNP — never
    assert _names(roles_client, kinds="distributor") == ["Опт"]
    assert _names(roles_client, kinds="manufacturer,dealer") == ["Завод", "Дилер"]
    assert _names(roles_client, kinds="unknown") == ["Ремонт", "Без роли"]  # rule said unknown + not determined
    assert _names(roles_client, only_manufacturers=True) == _names(roles_client, kinds="manufacturer") == ["Завод"]


def test_search_entity_and_min_score(roles_client):
    assert _names(roles_client, entity="ul") == ["Завод", "Дилер"]
    assert _names(roles_client, entity="ip") == ["Опт", "Ремонт"]  # «Без роли» has no INN: not guessed
    assert _names(roles_client, min_score=0.5) == ["Завод", "Опт", "Дилер"]  # 0.5 itself is kept
    assert _names(roles_client, kinds="manufacturer,distributor", entity="ip", min_score=0.6) == ["Опт"]


def test_search_entity_by_organizers_entity_type(search_client):
    """History: stats.entity_type of organizers' data; the row with a checksum-failing INN (inn=None) is still a ЮЛ."""
    def rows(**f):
        body = search_client.post("/api/search", json={"query_id": "q", "text": "бумага", "filters": f}).json()
        return [m["counterparty"]["inn"] for m in body["history"]]
    every = rows()
    ip = rows(entity="ip")
    assert sorted(ip) == ["470100000088", "780100000090"]
    assert rows(entity="ul") == [i for i in every if i not in ip] and None in rows(entity="ul")


def test_search_min_score_on_fixture(search_client):
    q = {"query_id": "q", "text": "бумага"}
    scores = [m["score"] for m in search_client.post("/api/search", json=q).json()["history"]]
    cut = scores[2]
    kept = search_client.post("/api/search", json={**q, "filters": {"min_score": cut}})
    got = [m["score"] for m in kept.json()["history"]]
    assert got == [s for s in scores if s >= cut] and len(got) >= 3 and int(kept.headers["X-Total-History"]) == len(got)


def test_search_facets_after_legal_status_each_group_after_the_other_groups(roles_client):
    """Liquidated / RNP never counted; a group is counted after the filters of the other groups (review of #240)."""
    q = {"query_id": "q", "text": "x", "filters": {"kinds": "distributor"}}
    r = roles_client.post("/api/search?limit=1", json=q)
    raw = r.headers["X-Facets"]
    assert raw.isascii()
    assert json.loads(raw) == {
        "history": {"kind": {"manufacturer": 1, "distributor": 1, "dealer": 1, "unknown": 2},  # kinds itself: off
                    "entity": {"ul": 0, "ip": 1, "unknown": 0},  # among distributors: «Опт», ИП
                    "near": {"worked": 1, "near": 1, "any": 1}},  # «Опт»: ИП, INN 78… (#242)
        "discovered": ZERO}
    # the page and the offset do not change the facets, nor does another choice in the same group
    again = roles_client.post("/api/search?limit=2&offset=1", json=q)
    assert json.loads(again.headers["X-Facets"]) == json.loads(raw)
    makers = roles_client.post("/api/search", json={**q, "filters": {"kinds": "manufacturer,dealer"}})
    assert json.loads(makers.headers["X-Facets"])["history"]["kind"] == json.loads(raw)["history"]["kind"]
    other = roles_client.post("/api/search", json={"query_id": "q", "text": "x",
                                                   "filters": {"entity": "ul", "min_score": 0.8}})
    assert json.loads(other.headers["X-Facets"])["history"] == {
        "kind": {"manufacturer": 1, "distributor": 0, "dealer": 0, "unknown": 0},  # ЮЛ with score ≥ 0.8: «Завод»
        "entity": {"ul": 1, "ip": 0, "unknown": 0},  # score ≥ 0.8: «Завод» (ЮЛ)
        "near": {"worked": 0, "near": 0, "any": 1}}  # «Завод»: ЮЛ without a KPP — region unknown


def test_search_facets_on_fixture(search_client):
    import json
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага", "filters": {"entity": "ip"}})
    facets = json.loads(r.headers["X-Facets"])
    assert facets["history"] == {"kind": {"manufacturer": 0, "distributor": 0, "dealer": 0, "unknown": 2},  # ИП
                                 "entity": {"ul": 4, "ip": 2, "unknown": 0},  # entity itself: off
                                 "near": {"worked": 2, "near": 2, "any": 2}}  # the two ИП: INN 78… / 47…, customers 78
    assert r.headers["X-Total-History"] == "2"


def test_facets_of_new_are_capped_at_the_block_size(search_client, monkeypatch):
    """A chip count is the block size when only that chip is chosen: «Новые» shows every candidate (#558)."""
    import json

    from src.enrich import discover as discover_mod
    makers = ["7801000012", "7801000020", "7801000037", "7801000044", "7801000051", "4701000038", "7804428656",
              "7805772436", "7807032363", "7810400592", "7811777270", "4703173460"]
    pool = [CounterpartyMatch(counterparty=Counterparty(inn=i, name=i, kind="manufacturer"), score=0.4,
                              origin="discovered", components=[]) for i in makers]
    pool.append(CounterpartyMatch(counterparty=Counterparty(inn="7707049388", name="опт", kind="distributor"),
                                  score=0.3, origin="discovered", components=[]))
    monkeypatch.setattr(discover_mod, "discover", lambda q, h, **k: discover_mod.Discovery(
        matches=pool[:discover_mod.MAX_SHOWN], candidates=pool, history_checked=True))
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert json.loads(r.headers["X-Facets"])["discovered"] == {
        "kind": {"manufacturer": len(makers), "distributor": 1, "dealer": 0, "unknown": 0},
        "entity": {"ul": len(makers) + 1, "ip": 0, "unknown": 0},
        "near": {"worked": 0, "near": 0, "any": len(makers) + 1}}  # no registry region, no KPP: not guessed


def test_facets_of_new_without_history_check_are_zero(search_client, monkeypatch):
    import json

    from src.enrich import discover as discover_mod
    m = CounterpartyMatch(counterparty=Counterparty(inn="7707049388", name="опт", kind="distributor"), score=0.3,
                          origin="discovered", components=[])
    monkeypatch.setattr(discover_mod, "discover", lambda q, h, **k: discover_mod.Discovery(
        matches=[m], candidates=[m], history_checked=False))
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert r.json()["discovered"] == [] and json.loads(r.headers["X-Facets"])["discovered"] == ZERO


def test_discovery_failure_is_a_header_and_history_survives(search_client, monkeypatch):
    """#287: a broken open source is told to the user (X-Discovery-Error), not shown as «no new companies»."""
    from src.enrich import discover as discover_mod

    def broken(q, h, **k):
        raise TimeoutError("/secret/path rmsp.nalog.ru: read timed out")

    monkeypatch.setattr(discover_mod, "discover", broken)
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert r.status_code == 200 and r.json()["history"]  # the history block is untouched
    assert r.json()["discovered"] == []
    assert r.headers["X-Discovery-Error"] == "TimeoutError"  # the type only: no message, no path

    class ОшибкаИсточника(RuntimeError):  # a non-ASCII type name must not make a latin-1 header fail (500)
        pass

    def broken_ru(q, h, **k):
        raise ОшибкаИсточника()

    monkeypatch.setattr(discover_mod, "discover", broken_ru)
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert r.status_code == 200 and r.headers["X-Discovery-Error"] == "Error"


def test_registry_refusal_inside_discovery_is_a_header(search_client, monkeypatch):
    """discover() catches a refused / failed registry request and goes on: still, the block may be incomplete."""
    from src.enrich import discover as discover_mod
    monkeypatch.setattr(discover_mod, "discover", lambda q, h, **k: discover_mod.Discovery(
        okpd_classes=["17.12"], history_checked=True, source_failed=True))
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert r.headers["X-Discovery-Error"] == app_mod.SOURCE_UNAVAILABLE


def test_discovery_without_history_check_is_a_header(search_client, monkeypatch):
    from src.enrich import discover as discover_mod
    m = CounterpartyMatch(counterparty=Counterparty(inn="7707049388", name="опт"), score=0.3, origin="discovered",
                          components=[])
    monkeypatch.setattr(discover_mod, "discover", lambda q, h, **k: discover_mod.Discovery(
        matches=[m], candidates=[m], okpd_classes=["17.12"], history_checked=False))
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert r.json()["discovered"] == [] and r.headers["X-Discovery-Error"] == app_mod.HISTORY_UNCHECKED


@pytest.mark.parametrize("found", [
    {"okpd_classes": ["17.12"], "history_checked": True},  # the sources answered: the block is what they gave
    {"okpd_classes": [], "history_checked": False},  # no OKPD2 class: nothing was searched, nothing failed
])
def test_no_discovery_header_when_nothing_failed(search_client, monkeypatch, found):
    from src.enrich import discover as discover_mod
    monkeypatch.setattr(discover_mod, "discover", lambda q, h, **k: discover_mod.Discovery(**found))
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert r.status_code == 200 and "X-Discovery-Error" not in r.headers


def test_corrected_query_is_a_header_and_as_typed_is_one_filter_away(search_client, monkeypatch, tmp_path):
    """#282: the search tells which text it used after fixing layout / typos; filters.spell=off — as typed."""
    from types import SimpleNamespace
    from urllib.parse import unquote

    from src.match import pipeline
    from src.match import search as search_mod
    from src.match.speller import Speller

    real = search_mod.load_config
    monkeypatch.setattr(search_mod, "load_config", lambda *a, **k: {**real(*a, **k), "spell": True})
    monkeypatch.setattr(pipeline, "get_index", lambda store, as_of: SimpleNamespace(path=tmp_path))
    monkeypatch.setattr(pipeline.speller, "for_index", lambda path: Speller({"бумага": 500}))
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумвга"})
    assert r.status_code == 200 and json.loads(unquote(r.headers["X-Query-Corrected"])) == [["бумвга", "бумага"]]
    assert r.json()["query"]["text"] == "бумвга"  # the answer keeps the query as the user typed it
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумвга", "filters": {"spell": "off"}})
    assert r.status_code == 200 and "X-Query-Corrected" not in r.headers
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert "X-Query-Corrected" not in r.headers  # nothing to fix — no header


@pytest.mark.parametrize("filters", [{"kinds": "producer"}, {"entity": "fl"}, {"min_score": "high"}])
def test_bad_filter_value_is_422_not_a_silent_no_filter(roles_client, filters):
    for path in ("/api/search", "/api/export"):
        r = roles_client.post(path, json={"query_id": "q", "text": "x", "filters": filters})
        assert r.status_code == 422 and "filters." in r.json()["detail"]


def _xlsx(client, **filters) -> list[dict]:
    r = client.post("/api/export", json={"query_id": "q", "text": "x", "filters": filters})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def test_export_role_and_entity_columns(roles_client):
    rows = _xlsx(roles_client)
    assert [(r["Наименование"], r["Роль"], r["ЮЛ/ИП"]) for r in rows] == [
        ("Завод", "Производитель", "ЮЛ"), ("Опт", "Дистрибьютор", "ИП/ФЛ"), ("Дилер", "Дилер", "ЮЛ"),
        ("Ремонт", "Роль не определена", "ИП/ФЛ"), ("Без роли", "Роль не определена", "не определено")]


def test_export_takes_the_same_filters(roles_client):
    assert [r["Наименование"] for r in _xlsx(roles_client, kinds="distributor,unknown")] == [
        "Опт", "Ремонт", "Без роли"]
    assert [r["Наименование"] for r in _xlsx(roles_client, entity="ip")] == ["Опт", "Ремонт"]
    assert [r["Наименование"] for r in _xlsx(roles_client, min_score=0.5)] == ["Завод", "Опт", "Дилер"]
    assert [r["Наименование"] for r in _xlsx(roles_client, only_manufacturers=True)] == ["Завод"]


# --- #242: «Может работать у заказчика» (near_customer), X-Search-Context, X-Facets.near ---------------------------

NEAR_API_ROWS = [  # name, inn, kpp, region (registry), customers_spb
    ("СПб", "7801000012", "780101001", None, 2),
    ("ЛО", "4701000038", "470101001", None, 0),
    ("Москва, работал в СПб", "7707049388", "770701001", None, 5),
    ("Москва", "7804428656", "770401001", None, 0),
    ("ИП СПб", "780100000090", None, None, 1),
    ("Регион неизвестен", None, None, None, None),
]


def _near_rows() -> list[CounterpartyMatch]:
    return [CounterpartyMatch(
        counterparty=Counterparty(inn=inn, name=name, kpp=kpp, region=region,
                                  stats={} if spb is None else {"customers_spb": spb}),
        score=0.5, origin="history",
        components=[ScoreComponent(name="product_fit", value=0.5, weight=1, explanation="x")])
        for name, inn, kpp, region, spb in NEAR_API_ROWS]


@pytest.fixture
def near_client(client, monkeypatch):
    """The search answers NEAR_API_ROWS; stub path — no cache, no discovery: the filter and the headers under test."""
    from src.match import search as search_mod
    monkeypatch.setattr(search_mod, "IS_STUB", True)
    monkeypatch.setattr(search_mod, "search", lambda q, k=20: _near_rows())
    return client


def _near_search(client, **q):
    r = client.post("/api/search", json={"query_id": "q", "text": "x", **q})
    assert r.status_code == 200, r.text
    return r, [m["counterparty"]["name"] for m in r.json()["history"]]


def test_search_near_customer_choices(near_client):
    every = [n for n, *_ in NEAR_API_ROWS]
    assert _near_search(near_client)[1] == every
    assert _near_search(near_client, filters={"near_customer": "any"})[1] == every
    assert _near_search(near_client, filters={"near_customer": "near"})[1] == ["СПб", "ЛО", "ИП СПб"]
    assert _near_search(near_client, filters={"near_customer": "worked"})[1] == [
        "СПб", "Москва, работал в СПб", "ИП СПб"]
    # the customer of the Leningrad region (INN 47): «worked» is its own region only
    lo = _near_search(near_client, customer_inn="4703173460", filters={"near_customer": "worked"})
    assert lo[1] == ["ЛО"] and lo[0].headers["X-Total-History"] == "1"


@pytest.mark.parametrize("okpd2, applied", [(["43.39.11.000"], "near"), (["17.12.14.110"], "any")])
def test_search_near_customer_auto(near_client, okpd2, applied):
    r, names = _near_search(near_client, okpd2=okpd2, filters={"near_customer": "auto"})
    assert names == _near_search(near_client, okpd2=okpd2, filters={"near_customer": applied})[1]
    assert json.loads(r.headers["X-Search-Context"])["applied_near"] == applied


def test_search_context_header(near_client):
    codes = ["81.21.10.000", "17.12.14.110", "81.22.11.000", "81.29.11.000", "81.10.10.000"]
    r, _ = _near_search(near_client, okpd2=codes, customer_inn="7804105239", filters={"near_customer": "auto"})
    raw = r.headers["X-Search-Context"]
    assert raw.isascii() and ", " not in raw and '": ' not in raw  # compact
    assert json.loads(raw) == {
        "customer_region": "Санкт-Петербург", "customer_basis": "ИНН 78…",
        "purchase_type": "services", "purchase_label": "услуга",
        "type_basis": ["81.21.10.000", "81.22.11.000", "81.29.11.000"],
        "default_near": "near", "applied_near": "near"}


def test_search_context_without_customer_inn_or_okpd2(near_client):
    r, _ = _near_search(near_client, region="Ленинградская область")
    assert json.loads(r.headers["X-Search-Context"]) == {
        "customer_region": "Ленинградская область", "customer_basis": "регион запроса",
        "purchase_type": None, "purchase_label": None, "type_basis": [],
        "default_near": "any", "applied_near": "any"}
    # an INN of another region (77, a federal body) does not move the customer: the query region stays
    r, _ = _near_search(near_client, customer_inn="7707049388", okpd2=["17.12.14.110"],
                        filters={"near_customer": "worked"})
    ctx = json.loads(r.headers["X-Search-Context"])
    assert (ctx["customer_region"], ctx["customer_basis"]) == ("Санкт-Петербург", "регион запроса")
    assert (ctx["purchase_type"], ctx["default_near"], ctx["applied_near"]) == ("goods", "any", "worked")


def test_search_facets_near_after_the_other_filters(near_client):
    r, _ = _near_search(near_client, filters={"near_customer": "near", "entity": "ip"})
    facets = json.loads(r.headers["X-Facets"])
    assert facets["history"] == {"kind": {"manufacturer": 0, "distributor": 0, "dealer": 0, "unknown": 1},
                                 "entity": {"ul": 2, "ip": 1, "unknown": 0},  # near: СПб, ЛО, ИП СПб
                                 "near": {"worked": 1, "near": 1, "any": 1}}  # ИП: «ИП СПб»
    assert facets["discovered"]["near"] == {"worked": 0, "near": 0, "any": 0}
    every = json.loads(_near_search(near_client)[0].headers["X-Facets"])["history"]["near"]
    assert every == {"worked": 3, "near": 3, "any": len(NEAR_API_ROWS)}
    lo = json.loads(_near_search(near_client, customer_inn="4703173460")[0].headers["X-Facets"])
    assert lo["history"]["near"] == {"worked": 1, "near": 3, "any": len(NEAR_API_ROWS)}


@pytest.mark.parametrize("q", [{"filters": {"near_customer": "close"}},
                               {"region": "Москва", "filters": {"near_customer": "near"}}])
def test_bad_near_customer_is_422(near_client, q):
    for path in ("/api/search", "/api/export"):
        r = near_client.post(path, json={"query_id": "q", "text": "x", **q})
        assert r.status_code == 422 and "filters.near_customer" in r.json()["detail"]


def test_export_takes_near_customer(near_client):
    assert [r["Наименование"] for r in _xlsx(near_client, near_customer="near")] == ["СПб", "ЛО", "ИП СПб"]
    r = near_client.post("/api/export", json={"query_id": "q", "text": "x", "okpd2": ["43.39.11.000"],
                                              "filters": {"near_customer": "auto"}})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    assert [row[1] for row in ws.iter_rows(min_row=2, values_only=True)] == ["СПб", "ЛО", "ИП СПб"]


def test_facets_near_of_new_count_the_candidates(search_client, monkeypatch):
    from src.enrich import discover as discover_mod
    pool = [CounterpartyMatch(counterparty=Counterparty(inn=i, name=i, kind="manufacturer", region=reg), score=0.4,
                              origin="discovered", components=[])
            for i, reg in (("7811777270", "Санкт-Петербург"), ("4703173460", "Ленинградская область"),
                           ("7810400592", "регион 77"), ("7807032363", None))]
    pool.append(CounterpartyMatch(  # liquidated: never counted, whatever discover() returns
        counterparty=Counterparty(inn="7805772436", name="ликв", region="Санкт-Петербург",
                                  status="В процессе ликвидации"), score=0.9, origin="discovered", components=[]))
    monkeypatch.setattr(discover_mod, "discover", lambda q, h, **k: discover_mod.Discovery(
        matches=pool, candidates=pool, history_checked=True))
    r = search_client.post("/api/search", json={"query_id": "q", "text": "бумага"})
    assert json.loads(r.headers["X-Facets"])["discovered"]["near"] == {"worked": 1, "near": 2, "any": 4}


def test_live_search_not_configured_is_a_503_with_the_reason(client, monkeypatch):
    """#550: without the Yandex key or the human's ceilings the button says so instead of failing."""
    monkeypatch.delenv("YANDEX_SEARCH_API_KEY", raising=False)
    r = client.post("/api/live-search", json={"query_id": "q", "text": "вакцины", "okpd2": ["21.20.10.110"]})
    assert r.status_code == 503 and r.json()["detail"].startswith("живой поиск не настроен")


def test_live_search_returns_the_found_companies_and_the_note(client, monkeypatch):
    from urllib.parse import unquote

    from src.enrich import discover as discover_mod
    found = CounterpartyMatch(counterparty=Counterparty(inn="7805772436", name="ООО"), score=0.5,
                              origin="discovered", components=[])
    monkeypatch.setattr(discover_mod, "live_search", lambda q: discover_mod.Live(matches=[found], note="Яндекс: ok"))
    monkeypatch.setattr(app_mod, "_verified", lambda ms: ms)  # the Evidence check has its own tests
    r = client.post("/api/live-search", json={"query_id": "q", "text": "вакцины", "okpd2": ["21.20.10.110"]})
    assert r.status_code == 200 and [m["counterparty"]["inn"] for m in r.json()] == ["7805772436"]
    assert unquote(r.headers["X-Live-Note"]) == "Яндекс: ok"
