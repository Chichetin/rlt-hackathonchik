"""The customer's list of purchases (#501): a template to fill in Excel -> the organizers' CSV of one test set.

`template()` is the empty workbook: «Закупки» (one row per purchase) and «Инструкция». `parse()` reads it back
(.xlsx, or a CSV saved from Excel: UTF-8 or cp1251, «;» / «,» / tab) with Russian synonyms of the column names
and returns Извещения.csv and ТРУ.csv text for src.ingest.testset plus the rows that did not pass and why.
A row is dropped only when there is nothing to search by (no subject and no OKPD2) or its number repeats;
a bad INN, OKPD2 code or price is dropped from the row with a note, the row stays.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.ingest.testset import ITEM_COLUMNS, NOTICE_COLUMNS

SHEET = "Закупки"
MAX_ROWS = 10_000
MAX_COLS = 50
MAX_UNZIPPED = 200 * 1024 * 1024  # a small .xlsx may unpack to gigabytes (zip bomb)
HEADERS = ("Номер закупки", "Предмет закупки", "ОКПД2", "ИНН заказчика", "НМЦК, руб.")
FIELDS = ("reqnum", "subject", "okpd2", "customer_inn", "start_price")
# normalized header (casefold, ё -> е, letters, digits and № only) -> field
SYNONYMS = {
    "reqnum": ("номерзакупки", "номер", "№", "№закупки", "реестровыйномер", "номерпроцедуры", "номерлота",
               "номеризвещения", "reqnum"),
    "subject": ("предметзакупки", "предмет", "наименование", "наименованиезакупки", "объектзакупки",
                "наименованиеобъектазакупки", "предметконтракта", "subject"),
    "okpd2": ("окпд2", "окпд", "кодокпд2", "кодокпд", "кодыокпд2", "okpd2", "okpd2code"),
    "customer_inn": ("иннзаказчика", "инн", "customerinn"),
    "start_price": ("нмцкруб", "нмцк", "нмцкрублей", "начальнаяцена", "начальнаямаксимальнаяценаконтракта", "цена",
                    "startprice"),
}
_BY_NAME = {name: f for f, names in SYNONYMS.items() for name in names}
_OKPD2 = re.compile(r"\d{2}(\.\d{1,3}){0,4}")
INSTRUCTION = (
    ("Как заполнить", ""),
    ("1", f"Одна строка листа «{SHEET}» — одна закупка. Заголовки не меняйте и не удаляйте."),
    ("2", "Обязателен предмет закупки или ОКПД2 (лучше оба): без них искать не по чему, строка пропускается."),
    ("3", "ОКПД2 — код с точками, например 17.12.14.110; несколько кодов — через запятую или «;»."),
    ("4", "ИНН заказчика (10 или 12 цифр) и НМЦК — по желанию: с ними порядок точнее (история заказчика, цена)."),
    ("5", "Номер закупки — по желанию, по нему строку найдёте в результате; повторять номер нельзя."),
    ("6", "Можно сохранить как .xlsx или как CSV (UTF-8 или Windows-1251, разделитель «;»)."),
    ("Результат", "Выгрузка «Excel» — лист «Рекомендации» по каждой закупке и лист «Ошибки»: строки, которые "
                  "не приняты, и замечания по принятым."),
)


class TemplateError(ValueError):
    """The file cannot be read as a list of purchases: the message says what to do."""


@dataclass
class RowError:
    row: int  # the row of the user's file, 1-based, the header is row 1
    reqnum: str
    reason: str
    kept: bool  # True: the row is searched, without the bad value


@dataclass
class Parsed:
    notices: str  # Извещения.csv text
    items: str  # ТРУ.csv text
    lots: int
    errors: list[RowError] = field(default_factory=list)


def template() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = SHEET
    ws.append(list(HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDE", (22, 60, 22, 16, 16), strict=True):
        ws.column_dimensions[col].width = width
        for row in range(2, 1002):  # text: Excel must not turn 0172…, 17.20 or an INN into a number
            ws[f"{col}{row}"].number_format = "@"
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Инструкция")
    for line in INSTRUCTION:
        info.append(list(line))
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm(header: object) -> str:
    return re.sub(r"[^0-9a-zа-я№]", "", str(header or "").casefold().replace("ё", "е"))


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _xlsx_rows(data: bytes) -> list[list[str]]:
    unreadable = "файл не читается как .xlsx — сохраните его в Excel заново или как CSV"
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                raise TemplateError("файл .xlsx слишком большой в распакованном виде — оставьте только лист закупок")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        if isinstance(e, TemplateError):
            raise
        raise TemplateError(unreadable) from e
    rows: list[list[str]] = []
    try:
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        for values in ws.iter_rows(max_col=MAX_COLS, values_only=True):
            rows.append([_cell(v) for v in values])
            if len(rows) > MAX_ROWS + 1:
                break
    except Exception as e:  # broken sheet XML (ParseError), a book without sheets: the user's file, not our bug
        raise TemplateError(unreadable) from e
    finally:
        wb.close()
    return rows


def _csv_rows(data: bytes) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1251", errors="replace")  # Excel «CSV (разделители - запятые)» in a Russian Windows
    lines = text.splitlines()
    head = lines[0] if lines else ""
    delimiter = max(";\t,", key=head.count)
    return [[c.strip() for c in r] for r in csv.reader(lines, delimiter=delimiter, quotechar='"')]


def _inn_ok(inn: str) -> bool:
    if not re.fullmatch(r"\d{10}|\d{12}", inn):
        return False
    d = [int(c) for c in inn]

    def check(weights: tuple[int, ...]) -> int:
        return sum(w * x for w, x in zip(weights, d, strict=False)) % 11 % 10

    if len(d) == 10:
        return check((2, 4, 10, 3, 5, 9, 4, 6, 8)) == d[9]
    return (check((7, 2, 4, 10, 3, 5, 9, 4, 6, 8)) == d[10]
            and check((3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8)) == d[11])


def _table(text_rows: list[list[str]]) -> str:
    buf = io.StringIO()
    csv.writer(buf, delimiter=";", lineterminator="\n").writerows(text_rows)
    return buf.getvalue()


def parse(name: str, data: bytes) -> Parsed:
    """The filled template (.xlsx or .csv) -> the organizers' CSV text; TemplateError when nothing is usable."""
    low = name.casefold()
    if low.endswith(".xlsx"):
        rows = _xlsx_rows(data)
    elif low.endswith(".csv"):
        rows = _csv_rows(data)
    else:
        raise TemplateError(f"{name}: нужен .xlsx или .csv по шаблону (скачайте шаблон и заполните его)")
    top = next((i for i, r in enumerate(rows) if any(r)), None)
    if top is None:
        raise TemplateError(f"{name}: пустой файл — нет строки заголовков")
    if len(rows) - top > MAX_ROWS + 1:
        raise TemplateError(f"{name}: больше {MAX_ROWS} строк — разбейте список на несколько файлов")
    cols: dict[str, int] = {}
    for i, h in enumerate(rows[top]):
        f = _BY_NAME.get(_norm(h))
        if f and f not in cols:
            cols[f] = i
    if "subject" not in cols and "okpd2" not in cols:
        raise TemplateError(f"{name}: нет колонок «Предмет закупки» и «ОКПД2» — заполните шаблон "
                            f"(колонки: {', '.join(HEADERS)})")

    notices: list[list[str]] = [list(NOTICE_COLUMNS)]
    items: list[list[str]] = [list(ITEM_COLUMNS)]
    errors: list[RowError] = []
    seen: dict[str, int] = {}
    for n, raw in enumerate(rows[top + 1:], start=top + 2):
        v = {f: (raw[i] if i < len(raw) else "") for f, i in cols.items()}
        if not any(v.values()):
            continue  # an empty row of the template
        reqnum = v.get("reqnum", "")

        def note(reason: str, kept: bool = True, row: int = n, num: str = reqnum) -> None:
            errors.append(RowError(row=row, reqnum=num, reason=reason, kept=kept))

        codes: list[str] = []
        for code in filter(None, (c.strip() for c in re.split(r"[,;\s]+", v.get("okpd2", "")))):
            if _OKPD2.fullmatch(code):
                codes.append(code)
            else:
                note(f"ОКПД2 «{code}» — не код вида 17.12.14.110, не учитывается")
        subject = v.get("subject", "")
        if not subject and not codes:
            note("нет ни предмета закупки, ни ОКПД2 — искать не по чему", kept=False)
            continue
        if reqnum and reqnum in seen:
            note(f"номер повторяется (уже в строке {seen[reqnum]}) — строка пропущена", kept=False)
            continue
        if reqnum:
            seen[reqnum] = n
        inn = v.get("customer_inn", "").replace(" ", "")
        if inn and not _inn_ok(inn):
            note(f"ИНН заказчика «{inn}» не проходит проверку контрольных цифр — не учитывается")
            inn = ""
        price = v.get("start_price", "").replace(" ", "").replace(" ", "").replace(",", ".")
        try:
            float(price) if price else None
        except ValueError:
            note(f"НМЦК «{v.get('start_price')}» не число — не учитывается")
            price = ""
        lot_id = str(n)  # the row of the file: digits, as /api/lot wants; the uploaded set is looked up first
        row = {"lot_id": lot_id, "reqnum": reqnum, "subject": subject, "customer_inn": inn, "start_price": price}
        notices.append([row.get(c, "") for c in NOTICE_COLUMNS])
        items.extend([lot_id, "", code] for code in codes)
    lots = len(notices) - 1
    if not lots:
        first = "; ".join(f"строка {e.row}: {e.reason}" for e in errors[:3])
        raise TemplateError(f"{name}: ни одной закупки для подбора" + (f" ({first})" if first else ""))
    return Parsed(notices=_table(notices), items=_table(items), lots=lots, errors=errors)
